"""Filial kassa qoldiqlari — alohida parser (ATM / filial reestri / kassirlarga ta'sir qilmaydi).

Excel ustunlari:
  Код БXM | Филиал (БХМ) номи | Сўм | Лимит сўм | Лимит доллар |
  АҚШ доллари 840 | Йена 392 | ... (valyuta summalari)
"""
from __future__ import annotations

import json
import logging
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.utils import get_column_letter

from .db import _connect

log = logging.getLogger(__name__)

# ISO 4217 numeric → meta. rate_uzs ≈ 1 birlik valyuta uchun so'm (ko'rsatish uchun).
# Kerak bo'lsa keyinroq Excel yoki API orqali yangilanadi.
CURRENCY_META: Dict[str, Dict[str, Any]] = {
    "840": {"code": "USD", "name": "АҚШ доллари", "name_ru": "Доллар США", "rate_uzs": 12800.0},
    "392": {"code": "JPY", "name": "Йена", "name_ru": "Японская иена", "rate_uzs": 85.0},
    "643": {"code": "RUB", "name": "Россия рубли", "name_ru": "Российский рубль", "rate_uzs": 140.0},
    "756": {"code": "CHF", "name": "Швейцария франки", "name_ru": "Швейцарский франк", "rate_uzs": 14500.0},
    "826": {"code": "GBP", "name": "Фунт стерлинг", "name_ru": "Фунт стерлингов", "rate_uzs": 16200.0},
    "978": {"code": "EUR", "name": "Евро", "name_ru": "Евро", "rate_uzs": 13800.0},
    "398": {"code": "KZT", "name": "Тенге", "name_ru": "Тенге", "rate_uzs": 26.0},
    "156": {"code": "CNY", "name": "Юан", "name_ru": "Юань", "rate_uzs": 1760.0},
    "972": {"code": "TJS", "name": "Сомони", "name_ru": "Сомони", "rate_uzs": 1180.0},
}

_FIXED_ALIASES = {
    "кодбхм": "bxm_code",
    "bxm": "bxm_code",
    "bxmcode": "bxm_code",
    "kodbxm": "bxm_code",
    "локалкод": "bxm_code",
    "localcode": "bxm_code",
    "филиалбхмноми": "branch_name",
    "филиалбхономи": "branch_name",
    "бхмноми": "branch_name",
    "бхономи": "branch_name",
    "филиалноми": "branch_name",
    "номи": "branch_name",
    "name": "branch_name",
    "сўм": "balance_uzs",
    "сум": "balance_uzs",
    "som": "balance_uzs",
    "uzs": "balance_uzs",
    "қолдиқ": "balance_uzs",
    "остаток": "balance_uzs",
    "лимитсўм": "limit_uzs",
    "лимитсум": "limit_uzs",
    "limituzs": "limit_uzs",
    "лимитдоллар": "limit_usd",
    "лимитдолларов": "limit_usd",
    "limitusd": "limit_usd",
}


def _norm_header(value: Any) -> str:
    if value is None:
        return ""
    s = str(value).strip().lower().replace("ё", "е")
    s = s.replace("ў", "у").replace("қ", "к").replace("ғ", "г").replace("ҳ", "х")
    s = re.sub(r"[\s_\-./\\()]+", "", s)
    return s


def _parse_amount(value: Any) -> Optional[float]:
    """Bo'sh / '-' → None (ma'lumot yo'q). 0 → 0.0."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip()
    if not s or s in ("-", "—", "–", "null", "None", ".", "n/a", "N/A"):
        return None
    s = s.replace("\u00a0", "").replace(" ", "").replace(",", ".")
    s = re.sub(r"[^\d.\-]", "", s)
    if not s or s in ("-", ".", "-."):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _resolve_header(norm: str, taken: set[str]) -> Optional[str]:
    if not norm:
        return None
    if norm in _FIXED_ALIASES:
        field = _FIXED_ALIASES[norm]
        return field if field not in taken else None

    # ISO numeric code in header: "акшдоллари840" / "евро978"
    m = re.search(r"(?<!\d)(840|392|643|756|826|978|398|156|972)(?!\d)", norm)
    if m:
        field = f"ccy_{m.group(1)}"
        return field if field not in taken else None

    if "лимит" in norm and ("долл" in norm or "usd" in norm or "840" in norm) and "limit_usd" not in taken:
        return "limit_usd"
    if "лимит" in norm and ("сум" in norm or "som" in norm or "uzs" in norm) and "limit_uzs" not in taken:
        return "limit_uzs"
    if norm in ("сум", "сом") or (norm.endswith("сум") and "лимит" not in norm):
        return "balance_uzs" if "balance_uzs" not in taken else None
    if ("код" in norm and ("бхм" in norm or "бхо" in norm or "bxm" in norm)) and "bxm_code" not in taken:
        return "bxm_code"
    if ("ном" in norm or "name" in norm) and ("бхм" in norm or "бхо" in norm or "филиал" in norm) and "branch_name" not in taken:
        return "branch_name"
    return None


def _match_headers(ws) -> Tuple[Dict[int, str], int]:
    header_row_idx = 1
    best_hits = 0
    for row_idx in range(1, min(8, (ws.max_row or 1) + 1)):
        row = next(ws.iter_rows(min_row=row_idx, max_row=row_idx, values_only=True))
        norms = [_norm_header(c) for c in row]
        hits = sum(1 for n in norms if n and _resolve_header(n, set()))
        if hits > best_hits:
            best_hits = hits
            header_row_idx = row_idx
        if hits >= 4:
            break

    column_map: Dict[int, str] = {}
    taken: set[str] = set()
    header_vals = next(ws.iter_rows(min_row=header_row_idx, max_row=header_row_idx, values_only=True))
    for col_idx, cell in enumerate(header_vals, start=1):
        field = _resolve_header(_norm_header(cell), taken)
        if field:
            column_map[col_idx] = field
            taken.add(field)
    return column_map, header_row_idx


def parse_branch_balances_xlsx(path: str | Path, sheet_name: Optional[str] = None) -> Dict[str, Any]:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Файл не найден: {path}")

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active
    column_map, header_row = _match_headers(ws)

    if "bxm_code" not in column_map.values() and "branch_name" not in column_map.values():
        header_vals = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
        detected = [
            f"{get_column_letter(i)}: '{c}'"
            for i, c in enumerate(header_vals, start=1)
            if c is not None
        ]
        wb.close()
        raise ValueError(
            "Не найдены колонки «Код БХМ» или «Филиал номи». "
            f"Строка заголовка #{header_row}: [{', '.join(detected[:20])}]"
        )

    records: List[Dict[str, Any]] = []
    errors: List[Dict[str, Any]] = []
    seen: set[str] = set()
    total_rows = 0

    for row_idx, row in enumerate(
        ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
    ):
        total_rows += 1
        if not row or all(c is None or (isinstance(c, str) and not str(c).strip()) for c in row):
            continue

        raw: Dict[str, Any] = {}
        for col_idx, field in column_map.items():
            value = row[col_idx - 1] if col_idx - 1 < len(row) else None
            raw[field] = value

        bxm = str(raw.get("bxm_code") or "").strip()
        # Excel ba'zan 11084.0 qilib beradi
        if re.fullmatch(r"\d+\.0+", bxm):
            bxm = bxm.split(".", 1)[0]
        name = str(raw.get("branch_name") or "").strip() or None

        if not bxm and not name:
            errors.append({"row": row_idx, "error": "Пустой код БХМ и название"})
            continue

        key = bxm or f"name:{name}"
        if key in seen:
            errors.append({"row": row_idx, "bxm_code": bxm, "error": "Дубликат в файле"})
            continue
        seen.add(key)

        currencies: Dict[str, Optional[float]] = {}
        for iso in CURRENCY_META:
            currencies[iso] = _parse_amount(raw.get(f"ccy_{iso}"))

        records.append({
            "bxm_code": bxm or None,
            "branch_name": name,
            "balance_uzs": _parse_amount(raw.get("balance_uzs")),
            "limit_uzs": _parse_amount(raw.get("limit_uzs")),
            "limit_usd": _parse_amount(raw.get("limit_usd")),
            "currencies": currencies,
        })

    wb.close()
    return {
        "records": records,
        "errors": errors,
        "total_rows": total_rows,
        "header_row": header_row,
        "columns": {str(k): v for k, v in column_map.items()},
    }


def init_branch_balance_tables() -> None:
    with _connect() as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS branch_balances (
                id              INTEGER PRIMARY KEY AUTOINCREMENT,
                bxm_code        TEXT,
                branch_name     TEXT,
                balance_uzs     REAL,
                limit_uzs       REAL,
                limit_usd       REAL,
                currencies_json TEXT NOT NULL DEFAULT '{}',
                matched_local_code TEXT,
                match_method    TEXT,
                imported_at     TEXT NOT NULL DEFAULT (datetime('now'))
            )
        """)
        conn.execute(
            "CREATE INDEX IF NOT EXISTS idx_branch_bal_bxm ON branch_balances(bxm_code)"
        )
        conn.execute(
            "CREATE INDEX IF NOT EXISTS idx_branch_bal_match ON branch_balances(matched_local_code)"
        )


def _norm_name(s: str) -> str:
    s = (s or "").strip().lower().replace("ё", "е")
    s = s.replace("ў", "у").replace("қ", "к").replace("ғ", "г").replace("ҳ", "х")
    return re.sub(r"\s+", " ", s)


def _match_branch(
    bxm_code: Optional[str],
    branch_name: Optional[str],
    branches_by_code: Dict[str, Dict[str, Any]],
    branches_by_name: List[Tuple[str, str]],
) -> Tuple[Optional[str], Optional[str]]:
    if bxm_code and bxm_code in branches_by_code:
        return bxm_code, "local_code"
    # ba'zan kod oldida 0
    if bxm_code:
        stripped = bxm_code.lstrip("0") or bxm_code
        for code in branches_by_code:
            if code.lstrip("0") == stripped:
                return code, "local_code"
    if branch_name:
        target = _norm_name(branch_name)
        for code, n in branches_by_name:
            if not n:
                continue
            if target == n or target in n or n in target:
                return code, "name"
    return None, None


def replace_branch_balances(records: List[Dict[str, Any]]) -> Dict[str, int]:
    """Eski qoldiqlarni o'chirib, yangilarini yozadi va filiallarga bog'laydi."""
    init_branch_balance_tables()
    with _connect() as conn:
        branches = [dict(r) for r in conn.execute("SELECT local_code, address, region FROM branches").fetchall()]
        by_code = {str(b["local_code"]): b for b in branches if b.get("local_code")}
        by_name: List[Tuple[str, str]] = []
        for b in branches:
            code = str(b.get("local_code") or "")
            # address ko'pincha "Номи, манзил" ko'rinishida
            addr = str(b.get("address") or "")
            name_part = addr.split(",")[0].strip() if addr else ""
            by_name.append((code, _norm_name(name_part)))
            by_name.append((code, _norm_name(addr)))

        conn.execute("DELETE FROM branch_balances")
        conn.execute("DELETE FROM sqlite_sequence WHERE name = 'branch_balances'")

        matched = 0
        unmatched = 0
        for rec in records:
            local, method = _match_branch(
                rec.get("bxm_code"), rec.get("branch_name"), by_code, by_name
            )
            if local:
                matched += 1
            else:
                unmatched += 1
            conn.execute(
                """
                INSERT INTO branch_balances (
                    bxm_code, branch_name, balance_uzs, limit_uzs, limit_usd,
                    currencies_json, matched_local_code, match_method
                ) VALUES (?,?,?,?,?,?,?,?)
                """,
                (
                    rec.get("bxm_code"),
                    rec.get("branch_name"),
                    rec.get("balance_uzs"),
                    rec.get("limit_uzs"),
                    rec.get("limit_usd"),
                    json.dumps(rec.get("currencies") or {}, ensure_ascii=False),
                    local,
                    method,
                ),
            )
    return {
        "saved": len(records),
        "matched_to_branches": matched,
        "unmatched": unmatched,
    }


def _enrich_balance_row(row: Dict[str, Any]) -> Dict[str, Any]:
    currencies_raw = {}
    try:
        currencies_raw = json.loads(row.get("currencies_json") or "{}")
    except json.JSONDecodeError:
        currencies_raw = {}

    currencies = []
    for iso, meta in CURRENCY_META.items():
        amount = currencies_raw.get(iso)
        if amount is None and iso not in currencies_raw:
            # key yo'q yoki null
            amount = currencies_raw.get(iso)
        rate = meta["rate_uzs"]
        uzs_eq = None if amount is None else round(float(amount) * rate, 2)
        currencies.append({
            "iso": iso,
            "code": meta["code"],
            "name": meta["name"],
            "name_ru": meta["name_ru"],
            "amount": amount,
            "rate_uzs": rate,
            "uzs_equivalent": uzs_eq,
        })

    bal = row.get("balance_uzs")
    lim = row.get("limit_uzs")
    usage_pct = None
    if bal is not None and lim is not None and float(lim) > 0:
        usage_pct = round(100.0 * float(bal) / float(lim), 1)

    usd_raw = currencies_raw.get("840")
    usd_amount = float(usd_raw) if usd_raw is not None else None
    limit_usd = row.get("limit_usd")
    usd_below_min = None
    if usd_amount is not None and limit_usd is not None:
        usd_below_min = float(usd_amount) < float(limit_usd)

    return {
        "bxm_code": row.get("bxm_code"),
        "branch_name": row.get("branch_name"),
        "balance_uzs": row.get("balance_uzs"),
        "limit_uzs": row.get("limit_uzs"),
        "limit_usd": row.get("limit_usd"),
        "usage_pct": usage_pct,
        "usd_amount": usd_amount,
        "usd_below_minimum": usd_below_min,
        "usd_status": ("below" if usd_below_min else "ok") if usd_below_min is not None else None,
        "matched_local_code": row.get("matched_local_code"),
        "match_method": row.get("match_method"),
        "imported_at": row.get("imported_at"),
        "currencies": currencies,
    }


def list_branch_balances() -> List[Dict[str, Any]]:
    init_branch_balance_tables()
    with _connect() as conn:
        rows = [dict(r) for r in conn.execute(
            "SELECT * FROM branch_balances ORDER BY bxm_code, id"
        ).fetchall()]
    return [_enrich_balance_row(r) for r in rows]


def get_branch_balance_by_local_code(local_code: str) -> Optional[Dict[str, Any]]:
    init_branch_balance_tables()
    with _connect() as conn:
        row = conn.execute(
            """
            SELECT * FROM branch_balances
             WHERE matched_local_code = ? OR bxm_code = ?
             ORDER BY CASE WHEN matched_local_code = ? THEN 0 ELSE 1 END, id DESC
             LIMIT 1
            """,
            (local_code, local_code, local_code),
        ).fetchone()
    return _enrich_balance_row(dict(row)) if row else None


def balances_by_local_code_map() -> Dict[str, Dict[str, Any]]:
    """Xarita uchun: local_code → balance payload."""
    out: Dict[str, Dict[str, Any]] = {}
    for b in list_branch_balances():
        key = b.get("matched_local_code") or b.get("bxm_code")
        if key:
            out[str(key)] = b
    return out
