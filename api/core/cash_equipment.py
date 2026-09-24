"""Kassa jihozlari ro'yxati — alohida parser (filial reestriga tegmaydi).

Excel ustunlari:
  № | Таркибий бўлинма коди (локал код) | Таркибий бўлинма номи |
  Асосий восита номи | Асосий воситанинг инвентар рақами |
  Асосий воситанинг тоифа рақами (номер под категории) |
  Балансга олинган санаси | Фойдаланишга топширилган санаси |
  Баланс бўйича тиклаш қиймати | Қолдиқ суммаси

«локал код» в этом файле — код региона (2600 = весь Ташкент), а не филиала,
поэтому филиал определяется по «Таркибий бўлинма номи» ↔ название филиала
(начало адреса в реестре: «Бунёдкор БХМ. 170119 …»). Если код всё-таки
совпал с 5-значным local_code филиала — берём его.
"""
from __future__ import annotations

import re
from difflib import SequenceMatcher
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.utils import get_column_letter

from .branch_balance import _match_branch_by_bxm_code, _norm_header, _parse_amount
from .db import _connect

FIELDS = (
    "kind",
    "local_code",
    "unit_name",
    "asset_name",
    "inventory_number",
    "category",
    "balance_date",
    "commissioned_date",
    "restoration_value",
    "residual_value",
)


def _resolve_header(norm: str, taken: set[str]) -> Optional[str]:
    if not norm:
        return None

    def pick(field: str) -> Optional[str]:
        return field if field not in taken else None

    # Tartib muhim: «тоифа рақами (номер под категории)» ichida «ном» ham bor.
    if "локал" in norm or ("булинма" in norm and "код" in norm):
        return pick("local_code")
    if "булинма" in norm and "ном" in norm:
        return pick("unit_name")
    if "тоифа" in norm or "категор" in norm:
        return pick("category")
    if "инвентар" in norm:
        return pick("inventory_number")
    if "восита" in norm and "ном" in norm:
        return pick("asset_name")
    if "фойдалан" in norm or "эксплуатац" in norm:
        return pick("commissioned_date")
    if "баланс" in norm and ("сана" in norm or "дата" in norm or "олин" in norm):
        return pick("balance_date")
    if "тиклаш" in norm or "восстанов" in norm:
        return pick("restoration_value")
    if "колдик" in norm or "остаточ" in norm:
        return pick("residual_value")
    return None


def _match_headers(ws) -> Tuple[Dict[int, str], int]:
    header_row_idx = 1
    best_hits = 0
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
        hits = sum(1 for c in row if _resolve_header(_norm_header(c), set()))
        if hits > best_hits:
            best_hits = hits
            header_row_idx = row_idx

    column_map: Dict[int, str] = {}
    taken: set[str] = set()
    header_vals = next(ws.iter_rows(min_row=header_row_idx, max_row=header_row_idx, values_only=True))
    for col_idx, cell in enumerate(header_vals, start=1):
        field = _resolve_header(_norm_header(cell), taken)
        if field:
            column_map[col_idx] = field
            taken.add(field)
    return column_map, header_row_idx


def _parse_date(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    s = str(value).strip()
    if not s or s in ("-", "—"):
        return None
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return s


def _clean_code(value: Any) -> str:
    s = str(value or "").strip()
    if re.fullmatch(r"\d+\.0+", s):
        s = s.split(".", 1)[0]
    return s


def _parse_sheet(ws, kind: str, records: List[Dict[str, Any]], errors: List[Dict[str, Any]]) -> int:
    """Один лист (Mashinka, Detektor …) → records. Возвращает число строк с техникой."""
    column_map, header_row = _match_headers(ws)
    if "asset_name" not in column_map.values() or not (
        {"local_code", "unit_name"} & set(column_map.values())
    ):
        return 0

    total = 0
    for row_idx, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        if not row or all(c is None or (isinstance(c, str) and not c.strip()) for c in row):
            continue
        raw = {f: (row[i - 1] if i - 1 < len(row) else None) for i, f in column_map.items()}

        asset = str(raw.get("asset_name") or "").strip()
        # Нумерация колонок под шапкой (1, 2, 3 …) и итоговые строки — пропускаем молча
        if not asset or asset.isdigit():
            continue
        total += 1
        code = _clean_code(raw.get("local_code")).upper() or None   # «27O00», «03B00», «1000»
        unit = str(raw.get("unit_name") or "").strip() or None
        if not code and not unit:
            errors.append({"row": row_idx, "sheet": kind, "asset_name": asset,
                           "error": "Нет ни кода, ни названия подразделения — строка пропущена"})
            continue

        records.append({
            "kind": kind,
            "local_code": code,
            "unit_name": unit,
            "asset_name": asset,
            "inventory_number": _clean_code(raw.get("inventory_number")) or None,
            "category": _clean_code(raw.get("category")) or None,
            "balance_date": _parse_date(raw.get("balance_date")),
            "commissioned_date": _parse_date(raw.get("commissioned_date")),
            "restoration_value": _parse_amount(raw.get("restoration_value")),
            "residual_value": _parse_amount(raw.get("residual_value")),
        })
    return total


def parse_cash_equipment_xlsx(path: str | Path, sheet_name: Optional[str] = None) -> Dict[str, Any]:
    """Читает все листы книги (Mashinka, Detektor, Termo printer …); имя листа = вид техники."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Файл не найден: {path}")

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets = [wb[sheet_name]] if sheet_name else list(wb.worksheets)

    records: List[Dict[str, Any]] = []
    errors: List[Dict[str, Any]] = []
    total_rows = 0
    by_sheet: Dict[str, int] = {}
    for ws in sheets:
        before = len(records)
        total_rows += _parse_sheet(ws, ws.title.strip(), records, errors)
        if len(records) > before:
            by_sheet[ws.title.strip()] = len(records) - before
    wb.close()

    if not records:
        raise ValueError(
            "Ни на одном листе не найдены колонки «Таркибий бўлинма номи» / «локал код» и «Асосий восита номи». "
            f"Листы: {', '.join(ws.title for ws in sheets)}"
        )
    return {
        "records": records,
        "errors": errors,
        "total_rows": total_rows,
        "by_sheet": by_sheet,
    }


def init_cash_equipment_tables() -> None:
    with _connect() as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS cash_equipment (
                id                 INTEGER PRIMARY KEY AUTOINCREMENT,
                local_code         TEXT,
                unit_name          TEXT,
                asset_name         TEXT,
                inventory_number   TEXT,
                category           TEXT,
                balance_date       TEXT,
                commissioned_date  TEXT,
                restoration_value  REAL,
                residual_value     REAL,
                matched_local_code TEXT,
                imported_at        TEXT NOT NULL DEFAULT (datetime('now'))
            )
        """)
        cols = {r["name"] for r in conn.execute("PRAGMA table_info(cash_equipment)").fetchall()}
        if "kind" not in cols:
            conn.execute("ALTER TABLE cash_equipment ADD COLUMN kind TEXT")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_cash_eq_match ON cash_equipment(matched_local_code)")


# Латиница, похожая на кириллицу («Xонобод», «Tошкент») → кириллица
_LOOKALIKE = str.maketrans("acekmhopxtyb", "асекмнорхтув")
# Служебные слова, которые пишут по-разному в двух файлах
_NAME_NOISE = re.compile(r"\b(бхм|бхо|минтакавий|маркази|марказ|банкинг|офис|ofis|офиси|филиали)\b")


def _name_key(name: Optional[str]) -> str:
    s = (name or "").lower().replace("ё", "е")
    s = s.replace("ў", "у").replace("қ", "к").replace("ғ", "г").replace("ҳ", "х")
    s = s.translate(_LOOKALIKE)
    s = re.sub(r"[\"'«»“”`]", " ", s)
    s = _NAME_NOISE.sub(" ", s)
    return re.sub(r"\s+", " ", s).strip()


def _branch_name_index(conn) -> Dict[str, str]:
    """name_key → local_code. Название филиала — начало адреса до первой точки."""
    index: Dict[str, str] = {}
    for b in conn.execute("SELECT local_code, address FROM branches").fetchall():
        if not b["local_code"] or not b["address"]:
            continue
        key = _name_key(str(b["address"]).split(".", 1)[0])
        if key:
            index.setdefault(key, str(b["local_code"]))
    return index


def _match_branch(rec: Dict[str, Any], codes: Dict[str, Any], names: Dict[str, str]) -> Optional[str]:
    code = str(rec.get("local_code") or "")
    if len(code) >= 5:
        local, _ = _match_branch_by_bxm_code(code, codes)
        if local:
            return local
    key = _name_key(rec.get("unit_name"))
    if not key:
        return None
    if key in names:
        return names[key]
    # «Premium» ↔ «Premium банкинг», «Урта бизнес» … — вхождение, если однозначно
    hits = {v for k, v in names.items() if k and (k in key or key in k) and min(len(k), len(key)) >= 4}
    if len(hits) == 1:
        return hits.pop()
    # Опечатки: «Чилонзар» ↔ «Чилонзор»
    scored = sorted(((SequenceMatcher(None, key, k).ratio(), v) for k, v in names.items()), reverse=True)
    if scored and scored[0][0] >= 0.85 and (len(scored) == 1 or scored[0][0] - scored[1][0] >= 0.05):
        return scored[0][1]
    return None


def _match_all(conn, records: List[Dict[str, Any]]) -> List[Optional[str]]:
    codes = {str(r["local_code"]).strip(): {} for r in conn.execute("SELECT local_code FROM branches").fetchall() if r["local_code"]}
    names = _branch_name_index(conn)
    return [_match_branch(rec, codes, names) for rec in records]


def replace_cash_equipment(records: List[Dict[str, Any]]) -> Dict[str, int]:
    """Eski ro'yxatni o'chirib, yangisini yozadi va filialga bog'laydi."""
    init_cash_equipment_tables()
    with _connect() as conn:
        matches = _match_all(conn, records)
        conn.execute("DELETE FROM cash_equipment")
        for rec, local in zip(records, matches):
            conn.execute(
                f"INSERT INTO cash_equipment ({', '.join(FIELDS)}, matched_local_code) "
                f"VALUES ({', '.join('?' * (len(FIELDS) + 1))})",
                (*[rec.get(f) for f in FIELDS], local),
            )
    matched = [m for m in matches if m]
    return {
        "saved": len(records),
        "matched": len(matched),
        "unmatched": len(records) - len(matched),
        "branches_matched": len(set(matched)),
    }


def rematch_cash_equipment() -> None:
    """Перепривязка уже загруженной техники (реестр филиалов мог обновиться)."""
    init_cash_equipment_tables()
    with _connect() as conn:
        rows = [dict(r) for r in conn.execute("SELECT id, local_code, unit_name, matched_local_code FROM cash_equipment").fetchall()]
        for row, local in zip(rows, _match_all(conn, rows)):
            if local != row["matched_local_code"]:
                conn.execute("UPDATE cash_equipment SET matched_local_code = ? WHERE id = ?", (local, row["id"]))


def clear_cash_equipment() -> int:
    init_cash_equipment_tables()
    with _connect() as conn:
        n = int(conn.execute("SELECT COUNT(*) AS n FROM cash_equipment").fetchone()["n"])
        conn.execute("DELETE FROM cash_equipment")
    return n


def list_cash_equipment(local_code: Optional[str] = None) -> List[Dict[str, Any]]:
    init_cash_equipment_tables()
    sql = "SELECT * FROM cash_equipment"
    params: Tuple[Any, ...] = ()
    if local_code:
        sql += " WHERE matched_local_code = ? OR local_code = ?"
        params = (local_code, local_code)
    sql += " ORDER BY id"
    with _connect() as conn:
        return [dict(r) for r in conn.execute(sql, params).fetchall()]


# ── Износ и срок замены ──────────────────────────────────────
# Колонки «когда менять» в Excel нет — считаем из амортизации:
#   износ = 1 − Қолдиқ / Тиклаш қиймати,
#   срок службы = лет в эксплуатации / износ (амортизация линейная),
#   дата полного списания = Фойдаланишга топширилган санаси + срок службы.
# По файлу срок у машинок выходит ~7 лет — он же берётся по умолчанию,
# когда посчитать нельзя (новая техника, нет суммы).
DEFAULT_LIFE_YEARS = 7.0
SOON_DAYS = 365


def equipment_lifecycle(item: Dict[str, Any], today: Optional[date] = None) -> Dict[str, Any]:
    today = today or date.today()
    restoration = item.get("restoration_value")
    residual = item.get("residual_value")
    try:
        start = date.fromisoformat(str(item.get("commissioned_date") or item.get("balance_date") or ""))
    except ValueError:
        start = None

    wear_pct: Optional[float] = None
    if restoration and restoration > 0 and residual is not None:
        wear_pct = round(max(0.0, min(1.0, 1 - float(residual) / float(restoration))) * 100, 1)

    writeoff: Optional[date] = None
    estimated = True
    if start:
        years = (today - start).days / 365.25
        if wear_pct is not None and 0 < wear_pct < 100 and years > 0.3 and wear_pct >= 2:
            writeoff = start + timedelta(days=round(years / (wear_pct / 100) * 365.25))
            estimated = False
        else:
            writeoff = start + timedelta(days=round(DEFAULT_LIFE_YEARS * 365.25))

    if residual is not None and float(residual) == 0 and restoration:
        status = "written_off"
        if writeoff and writeoff > today:
            writeoff = today
    elif writeoff and writeoff <= today:
        status = "written_off"
    elif writeoff and (writeoff - today).days <= SOON_DAYS:
        status = "soon"
    elif writeoff:
        status = "ok"
    else:
        status = "unknown"

    return {
        "wear_pct": wear_pct,
        "writeoff_date": writeoff.isoformat() if writeoff else None,
        "writeoff_estimated": estimated,
        "age_years": round((today - start).days / 365.25, 1) if start else None,
        "replace_status": status,
    }


_STATUS_ORDER = {"written_off": 0, "soon": 1, "ok": 2, "unknown": 3}


def equipment_replacement_plan(kind: Optional[str] = "Mashinka") -> Dict[str, Any]:
    """Список «что менять в первую очередь»: списанные (старые сверху), затем
    спишутся в ближайший год, затем остальные — по дате списания."""
    data = cash_equipment_by_branch()
    items: List[Dict[str, Any]] = []
    for g in data["branches"]:
        for it in g["items"]:
            if kind and it.get("kind") != kind:
                continue
            items.append({
                **it,
                "branch_local_code": g["local_code"],
                "branch_matched": g["matched"],
                "region": g.get("region"),
            })
    items.sort(key=lambda it: (_STATUS_ORDER.get(it["replace_status"], 9), it.get("writeoff_date") or "9999"))

    summary = {k: 0 for k in _STATUS_ORDER}
    for it in items:
        summary[it["replace_status"]] = summary.get(it["replace_status"], 0) + 1

    by_branch: Dict[str, Dict[str, Any]] = {}
    for it in items:
        b = by_branch.setdefault(it["branch_local_code"], {
            "local_code": it["branch_local_code"], "unit_name": it.get("unit_name"),
            "region": it.get("region"), "total": 0, "written_off": 0, "soon": 0,
        })
        b["total"] += 1
        if it["replace_status"] in ("written_off", "soon"):
            b[it["replace_status"]] += 1
    branches = sorted(by_branch.values(), key=lambda b: (-(b["written_off"] + b["soon"]), -b["written_off"]))

    return {
        "kind": kind,
        "today": date.today().isoformat(),
        "default_life_years": DEFAULT_LIFE_YEARS,
        "summary": summary,
        "total": len(items),
        "items": items,
        "branches": branches,
    }


def cash_equipment_by_branch() -> Dict[str, Any]:
    """Filial bo'yicha guruhlangan jihozlar + filial koordinatalari (xarita uchun)."""
    rematch_cash_equipment()
    rows = list_cash_equipment()
    today = date.today()
    with _connect() as conn:
        branches = {
            str(b["local_code"]): dict(b)
            for b in conn.execute("SELECT local_code, number, region, address, lat, lon FROM branches").fetchall()
        }

    groups: Dict[str, Dict[str, Any]] = {}
    for r in rows:
        matched = r.get("matched_local_code")
        key = matched or f"{r.get('local_code')}|{r.get('unit_name')}"
        g = groups.get(key)
        if g is None:
            b = branches.get(str(matched or "")) or {}
            g = groups[key] = {
                "local_code": matched or r.get("local_code"),
                "file_code": r.get("local_code"),
                "matched": bool(r.get("matched_local_code")),
                "unit_name": r.get("unit_name"),
                "number": b.get("number"),
                "region": b.get("region"),
                "address": b.get("address"),
                "lat": b.get("lat"),
                "lon": b.get("lon"),
                "count": 0,
                "restoration_total": 0.0,
                "residual_total": 0.0,
                "machines_written_off": 0,
                "machines_soon": 0,
                "items": [],
            }
        g["count"] += 1
        g["restoration_total"] += float(r.get("restoration_value") or 0)
        g["residual_total"] += float(r.get("residual_value") or 0)
        item = {k: r.get(k) for k in FIELDS}
        item.update(equipment_lifecycle(item, today))
        if item["kind"] == "Mashinka" and item["replace_status"] == "written_off":
            g["machines_written_off"] += 1
        elif item["kind"] == "Mashinka" and item["replace_status"] == "soon":
            g["machines_soon"] += 1
        g["items"].append(item)

    out = sorted(groups.values(), key=lambda g: g["local_code"] or "")
    return {
        "branches": out,
        "total_items": len(rows),
        "branches_count": len(out),
        "unmatched_branches": [f"{g['file_code']} · {g['unit_name']}" for g in out if not g["matched"]],
    }
