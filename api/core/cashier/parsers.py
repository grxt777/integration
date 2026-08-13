"""Excel and CSV report parsers for cashier KPI reports and HR staff registries."""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Dict, List, Tuple

import openpyxl

from .helpers import (
    BACK_KEYWORDS,
    SUMMARY_KEYWORDS,
    norm,
    num,
    parse_hr_status,
    sval,
)


def _classify_op_group(op_name: str) -> str:
    """Classify an operation name into 'back', 'front', or 'summary'."""
    n = norm(op_name)
    if any(kw in n for kw in SUMMARY_KEYWORDS):
        return 'summary'
    if any(kw in n for kw in BACK_KEYWORDS):
        return 'back'
    return 'front'


def _find_header_rows(all_rows: list) -> Tuple[int, int, int]:
    """
    Scan the first 30 rows to find the main header row index (0-based),
    optional sub-header row with 'Сони'/'минутда', and the data start row index.
    Returns (header_idx, sub_header_idx_or_-1, data_start_idx).
    """
    header_idx = None
    for r_idx in range(min(30, len(all_rows))):
        row = all_rows[r_idx]
        row_text = norm(' '.join(str(c or '') for c in row))
        # Look for characteristic FIO/tab column headers
        if any(k in row_text for k in ('ф и ш', 'фиш', 'ф и о', 'фио', 'fish', 'fio', 'табел')):
            header_idx = r_idx
            break

    if header_idx is None:
        header_idx = 1  # fallback: second row

    # Check if the next row is a sub-header (Сони / минутда labels)
    sub_header_idx = -1
    if header_idx + 1 < len(all_rows):
        next_row = all_rows[header_idx + 1]
        next_text = norm(' '.join(str(c or '') for c in next_row))
        if any(w in next_text for w in ('сони', 'минут', 'count', 'minut')):
            sub_header_idx = header_idx + 1

    data_start_idx = sub_header_idx + 1 if sub_header_idx >= 0 else header_idx + 1
    # Skip a possible "totals" row (row 4 in Kassirlar.xlsx which has grand totals)
    if data_start_idx < len(all_rows):
        totals_row = all_rows[data_start_idx]
        non_empty = [v for v in totals_row if v not in (None, '')]
        # If the first non-empty cell is a large integer (>= 1000) it's a totals row
        first_val = non_empty[0] if non_empty else None
        if isinstance(first_val, (int, float)) and first_val >= 1000:
            data_start_idx += 1

    return header_idx, sub_header_idx, data_start_idx


def _build_op_column_map(all_rows: list, header_idx: int, sub_header_idx: int) -> Dict[str, Any]:
    """Build a mapping of operation columns from the multi-row header."""
    header_row = all_rows[header_idx]
    sub_row = all_rows[sub_header_idx] if sub_header_idx >= 0 else [None] * len(header_row)

    n_cols = max(len(header_row), len(sub_row))
    header_row = list(header_row) + [None] * (n_cols - len(header_row))
    sub_row = list(sub_row) + [None] * (n_cols - len(sub_row))

    cur_group = ''
    col_label = [''] * n_cols  # final label for each column
    col_is_soni = [False] * n_cols
    col_is_minut = [False] * n_cols

    for idx in range(n_cols):
        hv = sval(header_row[idx])
        sv = sval(sub_row[idx])
        ns = norm(sv)

        if hv:
            cur_group = hv
        col_label[idx] = cur_group

        if 'сони' in ns or 'count' in ns or ('soni' in ns and 'ish' not in ns):
            col_is_soni[idx] = True
        if 'минут' in ns or 'minut' in ns:
            col_is_minut[idx] = True

    fixed = {}
    for idx, hv in enumerate(header_row):
        hv_s = sval(hv)
        nh = norm(hv_s).replace(' ', '')
        if not nh:
            continue

        if '№' in hv_s or nh in ('', 'nomer', '№') and idx == 0:
            fixed.setdefault('num', idx)
        elif any(k in nh for k in ('табел', 'tabel')):
            fixed.setdefault('tab', idx)
        elif any(k in nh for k in ('лавозим', 'lavozim', 'должность', 'position')):
            fixed.setdefault('pos', idx)
        elif any(k in nh for k in ('ф и ш', 'фиш', 'ф и о', 'фио', 'fish', 'fio', 'f i sh')):
            fixed.setdefault('fio', idx)
        elif any(k in nh for k in ('изох', 'izox', 'примечание', 'note', 'статус')):
            fixed.setdefault('note', idx)
        elif any(k in nh for k in ('локалкод', 'мфо', 'mfo', 'кодфилиал', 'localcod')):
            fixed.setdefault('bcode', idx)
        elif any(k in nh for k in ('бхмноми', 'bhm', 'филиалноми', 'наименованиефилиал', 'bhmnomi')):
            fixed.setdefault('bname', idx)
        elif any(k in nh for k in ('интизомий', 'jazoturi', 'jazatur', 'жазотури', 'jazoning', 'intizomiy', 'жазо', 'jazo', 'тур')):
            fixed.setdefault('disc_type', idx)
        elif any(k in nh for k in ('сана', 'sana', 'буйрук', 'buyruq', 'приказ', 'дата')):
            fixed.setdefault('disc_date_order', idx)
        elif any(k in nh for k in ('сабаб', 'sabab', 'причина')):
            fixed.setdefault('disc_reason', idx)
        elif any(k in nh for k in ('юклама', 'yuklama', 'хажми', 'hazmi')):
            fixed.setdefault('load', idx)
        elif any(k in nh for k in ('ишкуни', 'ishkuni')) and 'ишлаган' not in nh and 'ishlagan' not in nh:
            fixed.setdefault('std_days', idx)
        elif any(k in nh for k in ('ишлагакуни', 'ишлаганкуни', 'ishlagankuni', 'ishlagan')):
            fixed.setdefault('days', idx)

    # Positional auto-resolution if discipline columns are between bname and load
    bname_idx = fixed.get('bname')
    load_idx = fixed.get('load')

    if bname_idx is not None and load_idx is not None and load_idx > bname_idx + 1:
        between_cols = list(range(bname_idx + 1, load_idx))
        if len(between_cols) >= 3:
            fixed['disc_type'] = between_cols[0]
            fixed['disc_date_order'] = between_cols[1]
            fixed['disc_reason'] = between_cols[2]

    # Fallback if dynamic detection fails
    fixed.setdefault('num', 0)
    fixed.setdefault('tab', 1)
    fixed.setdefault('pos', 2)
    fixed.setdefault('fio', 3)
    fixed.setdefault('note', 4)
    fixed.setdefault('bcode', 5)
    fixed.setdefault('bname', 6)

    if fixed.get('disc_type') is not None or (load_idx is not None and load_idx >= 10) or n_cols >= 13:
        fixed.setdefault('disc_type', 7)
        fixed.setdefault('disc_date_order', 8)
        fixed.setdefault('disc_reason', 9)
        fixed.setdefault('load', 10)
        fixed.setdefault('std_days', 11)
        fixed.setdefault('days', 12)
    else:
        fixed.setdefault('disc_type', None)
        fixed.setdefault('disc_date_order', None)
        fixed.setdefault('disc_reason', None)
        fixed.setdefault('load', 7)
        fixed.setdefault('std_days', 8)
        fixed.setdefault('days', 9)

    ops = []
    bek_total_cnt = None
    bek_total_min = None
    front_total_cnt = None
    front_total_min = None

    first_op_col = max(fixed['days'] + 1, 10)
    i = first_op_col
    while i < n_cols:
        label = col_label[i]
        n_label = norm(label)

        if any(k in n_label for k in ('жами бек офис', 'jami bek ofis', 'жами бек', 'jami bek')):
            if col_is_soni[i]:
                bek_total_cnt = i
            elif col_is_minut[i]:
                bek_total_min = i
            if i + 1 < n_cols and col_is_minut[i + 1] and norm(col_label[i + 1]) == n_label:
                if bek_total_cnt is None:
                    bek_total_cnt = i
                bek_total_min = i + 1
                i += 2
                continue
        elif any(k in n_label for k in ('жами фронт офис', 'jami front ofis', 'жами фронт', 'jami front')):
            if col_is_soni[i]:
                front_total_cnt = i
            elif col_is_minut[i]:
                front_total_min = i
            if i + 1 < n_cols and col_is_minut[i + 1] and norm(col_label[i + 1]) == n_label:
                if front_total_cnt is None:
                    front_total_cnt = i
                front_total_min = i + 1
                i += 2
                continue

        if label and col_is_soni[i] and not any(k in n_label for k in ('жами', 'jami')):
            cnt_col = i
            min_col = i + 1 if (i + 1 < n_cols and col_is_minut[i + 1]) else None
            group = _classify_op_group(label)
            ops.append({
                'name': label,
                'group': group,
                'cnt_col': cnt_col,
                'min_col': min_col,
            })
            i += 2 if min_col is not None else 1
            continue

        i += 1

    fixed['bek_total_cnt'] = bek_total_cnt
    fixed['bek_total_min'] = bek_total_min
    fixed['front_total_cnt'] = front_total_cnt
    fixed['front_total_min'] = front_total_min

    return {'fixed': fixed, 'ops': ops}


def parse_cashiers_xlsx(path: str | Path) -> dict:
    """Parse KPI Excel report. Returns dict with 'records' list and metadata."""
    path = Path(path)
    path_str = str(path)

    if path_str.lower().endswith('.csv'):
        import csv
        with open(path, 'r', encoding='utf-8-sig', errors='ignore') as f:
            all_rows = [tuple(row) for row in csv.reader(f)]
    else:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        all_rows = list(ws.iter_rows(values_only=True))
        wb.close()

    if not all_rows:
        raise ValueError('Excel-файл пуст.')

    header_idx, sub_header_idx, data_start_idx = _find_header_rows(all_rows)
    col_map = _build_op_column_map(all_rows, header_idx, sub_header_idx)
    fx = col_map['fixed']
    ops_def = col_map['ops']

    records = []
    errors = []
    SKIP_NAMES = frozenset({
        'жами', 'итого', 'total', 'всего', 'сони', 'минут',
        'фиш', 'фио', 'ф и ш', 'ф и о', 'жами сони',
    })
    SKIP_TABS = frozenset({'жами', 'итого', 'total', 'всего', 'номер', '№'})

    for rn in range(data_start_idx, len(all_rows)):
        row = all_rows[rn]
        if not row or not any(v not in (None, '') for v in row):
            continue

        def get(col_idx):
            if col_idx is None or col_idx >= len(row):
                return None
            return row[col_idx]

        name = sval(get(fx['fio']))
        tab_num = sval(get(fx['tab'])) or sval(get(fx['num']))
        pos = sval(get(fx['pos']))
        note = sval(get(fx['note'])) if fx.get('note') is not None else ''
        bcode = sval(get(fx['bcode'])) if fx.get('bcode') is not None else ''
        bname = sval(get(fx['bname'])) if fx.get('bname') is not None else ''

        n_name = norm(name).replace(' ', '')
        n_tab = norm(tab_num).replace(' ', '')
        if n_name in SKIP_NAMES or n_tab in SKIP_TABS:
            continue
        if not name or name in ('-', '—', 'None', 'null', 'nan'):
            continue
        if re.match(r'^\d+$', name):
            continue

        load_raw = num(get(fx['load']))
        if 0 < load_raw <= 1.0:
            load_pct = round(load_raw * 100, 1)
        else:
            load_pct = round(load_raw, 1)

        std_days = num(get(fx.get('std_days')))
        days_worked = num(get(fx['days'])) or std_days

        if bcode and bname:
            branch_str = f"{bcode} - {bname}"
        elif bname:
            branch_str = bname
        elif bcode:
            branch_str = bcode
        else:
            branch_str = ''

        hr = parse_hr_status(note, pos, name)

        metrics = {}
        bek_ops_cnt = 0.0
        bek_ops_min = 0.0
        front_ops_cnt = 0.0
        front_ops_min = 0.0

        for op in ops_def:
            cnt = num(get(op['cnt_col']))
            mins = num(get(op['min_col'])) if op['min_col'] is not None else 0.0
            op_name = op['name']
            metrics[op_name] = {'count': cnt, 'minutes': mins}
            if op['group'] == 'back':
                bek_ops_cnt += cnt
                bek_ops_min += mins
            elif op['group'] == 'front':
                front_ops_cnt += cnt
                front_ops_min += mins

        bek_cnt_raw = num(get(fx.get('bek_total_cnt')))
        bek_min_raw = num(get(fx.get('bek_total_min')))
        front_cnt_raw = num(get(fx.get('front_total_cnt')))
        front_min_raw = num(get(fx.get('front_total_min')))

        bek_count = bek_cnt_raw if bek_cnt_raw > 0 else bek_ops_cnt
        bek_minutes = bek_min_raw if bek_min_raw > 0 else bek_ops_min
        front_count = front_cnt_raw if front_cnt_raw > 0 else front_ops_cnt
        front_minutes = front_min_raw if front_min_raw > 0 else front_ops_min

        ops_count = bek_count + front_count
        ops_minutes = bek_minutes + front_minutes

        computed_load = 0.0
        load_denom = std_days or days_worked
        if load_denom > 0 and ops_minutes > 0:
            computed_load = round((ops_minutes / (load_denom * 480)) * 100, 1)
        if load_pct == 0 and computed_load:
            load_pct = computed_load
        load_difference = round(abs(load_pct - computed_load), 1) if load_pct and computed_load else 0.0

        disc_type = sval(get(fx.get('disc_type'))) if fx.get('disc_type') is not None else ''
        disc_date_order = sval(get(fx.get('disc_date_order'))) if fx.get('disc_date_order') is not None else ''
        disc_reason = sval(get(fx.get('disc_reason'))) if fx.get('disc_reason') is not None else ''

        rec = {
            'full_name': name,
            'tab_number': tab_num,
            'position': pos,
            'days_worked': days_worked,
            'std_days': std_days,
            'branch_name': branch_str,
            'raw_note': note,
            'hr_status_code': hr['status_code'],
            'hr_status_label': hr['status_label'],
            'has_replacement': hr['has_replacement'],
            'replacing_full_name': None,
            'replaced_by_full_name': None,
            'operations_count': ops_count,
            'operations_minutes': ops_minutes,
            'bek_count': bek_count,
            'bek_minutes': bek_minutes,
            'front_count': front_count,
            'front_minutes': front_minutes,
            'load_percent': load_pct,
            'load_difference': load_difference,
            'employee_number': sval(get(fx['num'])),
            'metrics': metrics,
            'discipline_type': disc_type or None,
            'discipline_date_order': disc_date_order or None,
            'discipline_reason': disc_reason or None,
        }
        records.append(rec)

    if not records:
        raise ValueError(
            f'После строки шапки {header_idx + 1} не найдено валидных строк кассиров.'
        )

    def _empty_seat(r: dict) -> bool:
        return r.get('hr_status_code') in ('maternity', 'vacation') and float(r.get('operations_count') or 0) <= 0

    # Adjacent correlation (вакт. directly above empty-seat декрет/мехнат.тат.)
    for i in range(len(records) - 1):
        r_curr = records[i]
        r_next = records[i + 1]
        if (r_curr['hr_status_code'] == 'temporary'
                and _empty_seat(r_next)
                and not r_curr.get('replacing_full_name')
                and not r_next.get('replaced_by_full_name')):
            r_curr['replacing_full_name'] = r_next['full_name']
            r_next['replaced_by_full_name'] = r_curr['full_name']
            r_next['has_replacement'] = 1
            r_curr['hr_status_label'] = f"🟡 Временный (замещает {r_next['full_name']})"
            if r_next['hr_status_code'] == 'maternity':
                r_next['hr_status_label'] = f"🟣 В декрете (замещает {r_curr['full_name']})"
            elif r_next['hr_status_code'] == 'vacation':
                r_next['hr_status_label'] = f"🔵 В отпуске (замещает {r_curr['full_name']})"

    # Filial-group correlation
    branch_groups: dict[str, list] = {}
    for r in records:
        b_key = r.get('branch_name') or 'Default'
        branch_groups.setdefault(b_key, []).append(r)

    for b_key, b_recs in branch_groups.items():
        unmatched_temps = [r for r in b_recs if r['hr_status_code'] == 'temporary' and not r.get('replacing_full_name')]
        unmatched_absents = [r for r in b_recs if _empty_seat(r) and not r.get('replaced_by_full_name')]
        for temp_r, abs_r in zip(unmatched_temps, unmatched_absents):
            temp_r['replacing_full_name'] = abs_r['full_name']
            abs_r['replaced_by_full_name'] = temp_r['full_name']
            abs_r['has_replacement'] = 1
            temp_r['hr_status_label'] = f"🟡 Временный (замещает {abs_r['full_name']})"
            if abs_r['hr_status_code'] == 'maternity':
                abs_r['hr_status_label'] = f"🟣 В декрете (замещает {temp_r['full_name']})"
            elif abs_r['hr_status_code'] == 'vacation':
                abs_r['hr_status_label'] = f"🔵 В отпуске (замещает {temp_r['full_name']})"

        for abs_r in b_recs:
            if not _empty_seat(abs_r) or abs_r.get('replaced_by_full_name'):
                continue
            abs_r['has_replacement'] = 0
            if abs_r['hr_status_code'] == 'maternity':
                abs_r['hr_status_label'] = '🟣 В декрете (без замены!)'
            elif abs_r['hr_status_code'] == 'vacation':
                abs_r['hr_status_label'] = '🔵 В отпуске (без замены!)'

    return {
        'records': records,
        'errors': errors,
        'header_rows': [header_idx + 1, sub_header_idx + 1 if sub_header_idx >= 0 else header_idx + 1],
        'columns': [op['name'] for op in ops_def],
    }


def parse_cashier_status_xlsx(path: str | Path) -> dict:
    """Parse the HR staff registry Excel (Штат.xlsx)."""
    path = Path(path)
    path_str = str(path)

    if path_str.lower().endswith('.csv'):
        import csv
        with open(path, 'r', encoding='utf-8-sig', errors='ignore') as f:
            all_rows = [tuple(row) for row in csv.reader(f)]
    else:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        all_rows = list(ws.iter_rows(values_only=True))
        wb.close()

    if not all_rows:
        raise ValueError('Excel-файл пуст.')

    fio_col = None
    pos_col = None
    mfo_col = None
    note_col = None
    header_row_idx = None

    for r_idx in range(min(30, len(all_rows))):
        row = all_rows[r_idx]
        if not row:
            continue
        r_vals = [sval(x) for x in row]
        n_vals = [norm(x) for x in r_vals]

        f_c = p_c = m_c = n_c = b_c = None
        for idx, nv in enumerate(n_vals):
            nv_nospace = nv.replace(' ', '')
            if any(w in nv_nospace for w in ('фиш', 'фио', 'fio', 'fish', 'сотрудник', 'работник', 'xodim', 'ф и ш', 'ф и о')):
                f_c = idx
            elif any(w in nv_nospace for w in ('таркибий', 'лавозим', 'должность', 'lavozim', 'position')) and \
                    not any(k in nv_nospace for k in ('маош', 'оклад', 'разряд', 'коэффициент')):
                p_c = idx
            elif any(w in nv_nospace for w in ('mfo', 'мфо', 'кодфилиала')):
                m_c = idx
            elif any(w in nv_nospace for w in ('филиал', 'бхм', 'бхо', 'подразделение', 'наименованиефилиала', 'bhm', 'branch')):
                b_c = idx
            elif any(w in nv_nospace for w in ('изох', 'примечание', 'статус', 'причина', 'note')):
                n_c = idx

        if f_c is not None or p_c is not None:
            fio_col, pos_col, mfo_col, note_col, bname_col = f_c, p_c, m_c, n_c, b_c
            header_row_idx = r_idx
            break

    if fio_col is None:
        fio_col = 7
    if pos_col is None:
        pos_col = 2
    if mfo_col is None:
        mfo_col = 1
    if note_col is None:
        note_col = 10
    bname_col = b_c if 'b_c' in locals() else None

    raw_parsed = []
    current_branch = 'Не указан'
    current_pos = ''
    current_mfo = ''

    start_idx = (header_row_idx + 1) if header_row_idx is not None else 0

    SKIP_FIO = frozenset({
        'ф и ш', 'ф и о', 'фио', 'фиш', 'jami', 'итого', 'сони', 'минут',
        'номер', '№', 'ставка', 'всего', 'жами',
    })

    for rn in range(start_idx, len(all_rows)):
        row = all_rows[rn]
        if not row or not any(v not in (None, '') for v in row):
            continue

        r_str = [sval(cell) for cell in row]

        def gcol(idx, r=r_str):
            return r[idx] if idx is not None and idx < len(r) else ''

        first_val = r_str[0] if r_str else ''
        n_first = norm(first_val)

        if any(w in n_first for w in ('регион', 'region')):
            continue
        if first_val.startswith('Филиал:') or any(w in n_first for w in ('филиал', 'бхм', 'бхо', 'центр', 'офис')):
            current_branch = first_val.replace('Филиал:', '').strip()
            current_mfo = gcol(mfo_col)
            continue
        if first_val.startswith('**'):
            current_pos = first_val.replace('**', '').strip()
            continue
        if any(w in n_first for w in ('жами:', 'жами', 'итого', 'total')) and len(first_val) < 20:
            continue

        fio_val = gcol(fio_col)
        pos_val = gcol(pos_col)
        mfo_val = gcol(mfo_col)
        bname_val = gcol(bname_col) if bname_col is not None else ''
        note_val = gcol(note_col)

        is_subrow = not pos_val.strip() or pos_val.strip() in ('-', '—')

        if not is_subrow:
            current_pos = pos_val

        if is_subrow:
            pos_val = current_pos

        if bname_val and bname_val not in ('-', '—') and len(bname_val) > 2:
            current_branch = bname_val

        if not mfo_val or mfo_val in ('-', '—'):
            mfo_val = current_mfo
        else:
            current_mfo = mfo_val

        if not fio_val:
            for test_idx in (7, 6, 5, 3, 2):
                val = r_str[test_idx] if len(r_str) > test_idx else ''
                n_v = norm(val)
                if val and len(val) >= 3 and n_v not in SKIP_FIO:
                    if any(w in n_v for w in ('вакант', 'вакансия', 'vakant', 'vakansiya', 'bosh', 'bush')) or len(val.split()) >= 2:
                        fio_val = val
                        break

        n_fio = norm(fio_val)
        if not fio_val or len(fio_val) < 3 or n_fio in SKIP_FIO:
            continue

        is_vacant = any(w in n_fio for w in ('вакант', 'вакансия', 'vakant', 'vakansiya', 'bosh', 'буш', 'bush', 'свобод'))
        full_name = f"ВАКАНСИЯ ({pos_val if pos_val else 'Кассир'})" if is_vacant else fio_val

        mfo_clean = mfo_val.replace('\xa0', '').strip()
        branch_label = current_branch
        if mfo_clean and mfo_clean not in current_branch:
            branch_label = f"{mfo_clean} - {current_branch}"

        clean_note = note_val.replace('\xa0', '').strip()
        if clean_note.isdigit():
            clean_note = ''

        raw_parsed.append((rn, is_subrow, {
            'branch_name': branch_label,
            'position': pos_val if pos_val else 'Кассир',
            'full_name': full_name,
            'raw_note': clean_note,
            'status_code': 'vacant' if is_vacant else 'active',
            'status_label': '⚪ Вакант (Свободная ставка)' if is_vacant else '🟢 Ишлаяпти',
            'replacing_full_name': None,
            'replaced_by_full_name': None,
            'has_replacement': 0,
        }))

    records = [rec for _, _, rec in raw_parsed]

    if not records:
        raise ValueError('Не найдено ни одной валидной строки с ФИО кассиров или вакансиями в реестре штата.')

    # Pass 1: Determine status_code from raw_note
    for r in records:
        if r['status_code'] == 'vacant':
            continue
        n_flat = norm(f"{r['position']} {r['raw_note']}").replace(' ', '')
        if 'dekret' in n_flat or 'декрет' in n_flat:
            r['status_code'] = 'maternity'
        elif any(w in n_flat for w in ('mexnattat', 'мехнатtat', 'мехнаттатил', 'mexnattatil', 'татил', 'tatil')):
            r['status_code'] = 'vacation'
        elif any(w in n_flat for w in ('mexnat', 'мехнат', 'отпуск')):
            r['status_code'] = 'vacation'
        elif any(w in n_flat for w in ('vakt', 'вакт', 'zamesh', 'замещ')):
            r['status_code'] = 'temporary'
        elif any(w in n_flat for w in ('kasal', 'касал', 'больн')):
            r['status_code'] = 'sick'

    # Pass 2: Sub-row adjacency pairing
    for i in range(len(raw_parsed) - 1):
        _, is_sub_curr, r_curr = raw_parsed[i]
        _, is_sub_next, r_next = raw_parsed[i + 1]

        if (r_curr['status_code'] == 'temporary'
                and is_sub_next
                and r_next['status_code'] in ('maternity', 'vacation')
                and not r_curr.get('replacing_full_name')
                and not r_next.get('replaced_by_full_name')):
            r_curr['replacing_full_name'] = r_next['full_name']
            r_next['replaced_by_full_name'] = r_curr['full_name']
            r_next['has_replacement'] = 1
        elif (is_sub_curr
                and r_curr['status_code'] in ('maternity', 'vacation')
                and r_next['status_code'] == 'temporary'
                and not r_curr.get('replaced_by_full_name')
                and not r_next.get('replacing_full_name')):
            r_next['replacing_full_name'] = r_curr['full_name']
            r_curr['replaced_by_full_name'] = r_next['full_name']
            r_curr['has_replacement'] = 1

    # Pass 3: Within same branch pairing
    branch_recs: dict[str, list] = {}
    for r in records:
        bk = r.get('branch_name') or 'Default'
        branch_recs.setdefault(bk, []).append(r)

    for bk, b_list in branch_recs.items():
        pos_groups: dict[str, list] = {}
        for r in b_list:
            pg = norm(r.get('position', '')).replace(' ', '')
            pos_groups.setdefault(pg, []).append(r)

        for pg, pg_list in pos_groups.items():
            unmatched_temps = [r for r in pg_list if r['status_code'] == 'temporary' and not r.get('replacing_full_name')]
            unmatched_absents = [r for r in pg_list if r['status_code'] in ('maternity', 'vacation') and not r.get('replaced_by_full_name')]
            for temp_r, abs_r in zip(unmatched_temps, unmatched_absents):
                temp_r['replacing_full_name'] = abs_r['full_name']
                abs_r['replaced_by_full_name'] = temp_r['full_name']
                abs_r['has_replacement'] = 1

        for r in b_list:
            if r['status_code'] in ('maternity', 'vacation') and not r.get('replaced_by_full_name'):
                r['has_replacement'] = 0

    # Pass 4: Build readable labels
    for r in records:
        sc = r['status_code']
        if sc == 'vacant':
            r['status_label'] = '⚪ Вакант (Свободная ставка)'
        elif sc == 'temporary':
            r['status_label'] = (
                f"🟡 Вақтинча (ўринбосар: {r['replacing_full_name']})"
                if r.get('replacing_full_name') else '🟡 Вақтинча ходим'
            )
        elif sc == 'maternity':
            r['status_label'] = (
                f"🟣 Декретда (ўрнида: {r['replaced_by_full_name']})"
                if r.get('replaced_by_full_name') else '🟣 Декретда (без замены!)'
            )
        elif sc == 'vacation':
            r['status_label'] = (
                f"🔵 Меҳнат татилда (ўрнида: {r['replaced_by_full_name']})"
                if r.get('replaced_by_full_name') else '🔵 Меҳнат татилда (без замены!)'
            )
        elif sc == 'sick':
            r['status_label'] = '🔴 Касал (Больничный)'
        else:
            r['status_code'] = 'active'
            r['status_label'] = '🟢 Ишлаяпти'
            r['has_replacement'] = 1

    return {'records': records, 'total_rows': len(all_rows)}
